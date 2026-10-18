"""Genera el reporte del Avance 3 en HTML, PDF y DOCX a partir de un solo contenido.

Las figuras se toman de la libreta ejecutada (Avance3.48.ipynb), de modo que
el reporte y la libreta siempre muestran los mismos resultados. Uso, desde la raíz del repositorio:
    python3 scripts/generar_documento_avance3.py
Requiere Google Chrome (PDF), PyMuPDF (pie de página e índice) y python-docx.
"""
import base64
import html
import io
import json
import pathlib
import re
import subprocess

import fitz
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

RAIZ = pathlib.Path(__file__).resolve().parent.parent
CARPETA = RAIZ
LIBRETA = CARPETA / "Avance3.48.ipynb"
CSS = CARPETA / "assets" / "print.css"
LOGO = CARPETA / "assets" / "logo_tec.png"
HTML = CARPETA / "Avance3.48.html"
PDF = CARPETA / "Avance3.48.pdf"
DOCX = CARPETA / "Avance3.48.docx"
CHROME = "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome"
REPO = "https://github.com/oscar-ramirez-anaya/Avance3_Equipo48"
FECHA = "18 de octubre de 2026"

AZUL = RGBColor(0x0B, 0x3B, 0x8C)
AZUL_TABLA = "1346BD"


# ------------------------------------------------------------------ figuras de la libreta
def figuras_libreta():
    nb = json.loads(LIBRETA.read_text(encoding="utf-8"))
    figs = {}
    for c in nb["cells"]:
        for o in c.get("outputs", []):
            if "data" in o and "image/png" in o["data"]:
                figs[c["id"]] = o["data"]["image/png"]
    return figs


# ------------------------------------------------------------------ contenido
# Bloques: ("h2", id, texto) ("h3", id, texto) ("p", html) ("ul", [html]) ("ol", [html])
#          ("tabla", id, titulo, encabezados, filas) ("fig", id, celda, titulo) ("nota", titulo, html)
B = []
h2 = lambda i, t: B.append(("h2", i, t))
h3 = lambda i, t: B.append(("h3", i, t))
p = lambda t: B.append(("p", t))
ul = lambda xs: B.append(("ul", xs))
ol = lambda xs: B.append(("ol", xs))
tabla = lambda i, t, e, f: B.append(("tabla", i, t, e, f))
fig = lambda i, c, t: B.append(("fig", i, c, t))
nota = lambda t, x: B.append(("nota", t, x))

h2("s1", "1. Introducción")
h3("s11", "1.1 Propósito")
p("Este documento reporta la fase de <b>Ingeniería del modelo</b> del proyecto, conforme a la metodología CRISP-ML(Q) (Visengeriyeva et al., 2023). Se construye un modelo de referencia (<i>baseline</i>) para reconocer el producto en recortes de anaquel, con dos propósitos: establecer las medidas de calidad del modelo y fijar el marco de comparación de los modelos más avanzados. El documento responde cinco preguntas: qué algoritmo sirve como baseline, qué características son importantes, si el modelo sub o sobreajusta, cuál es la métrica adecuada y cuál debe ser el desempeño mínimo.")
p(f"La libreta ejecutada que respalda cada cifra, <code>Avance3.48.ipynb</code>, está en el repositorio <code>{REPO}</code>.")
h3("s12", "1.2 Problema y datos")
p("El planograma exige saber qué producto ocupa cada posición del anaquel. Una posición se declara correcta, faltante o mal ubicada comparando el producto reconocido contra el prescrito. Por eso la tarea que se modela es la <b>clasificación multiclase de recortes de producto</b>. El plan inicial contemplaba como línea base un detector preentrenado sobre escenas completas. La revisión de los datos en el servidor de Bnext mostró que las escenas no tienen anotaciones (ni cajas delimitadoras ni planogramas de referencia), así que un detector no podría evaluarse con métricas; además, el servidor no tiene GPU. Los recortes sí tienen una etiqueta por carpeta y permiten medir el componente de reconocimiento.")
p("Bnext amplió el conjunto respecto al Avance 2. Se usan todos los recortes etiquetados disponibles (Tabla 1): 35 carpetas con código de producto y 206 carpetas de agrupamiento. No se cuenta con una tabla que asocie cada agrupamiento con un producto, así que se asume que cada carpeta es un producto distinto; este supuesto se revisa en la sección 5.")
tabla("t1", "Recortes de producto disponibles", ["Tipo de carpeta", "Clases", "Recortes", "Recortes por clase (mín. / mediana / máx.)", "Captura"], [
    ["Con código (SKU_###)", "35", "422", "10 / 11 / 24", "enero de 2026"],
    ["Agrupamiento (GRP_###)", "206", "2 465", "5 / 11 / 32", "junio de 2026"],
    ["<b>Total</b>", "<b>241</b>", "<b>2 887</b>", "5 / 11 / 32", ""],
])
fig("f1", "celda-005", "Imágenes por clase según el tipo de carpeta.")
nota("Manejo de datos", "Conforme al acuerdo con Bnext, la libreta se ejecutó dentro del servidor de la organización y las imágenes no salieron de su infraestructura. Este reporte y la libreta publicada contienen solo tablas y gráficas agregadas. Las clases se reemplazaron por identificadores anónimos (SKU_### y GRP_###).")
h3("s13", "1.3 Preparación y partición")
p("Se reutilizó la extracción del Avance 2: 822 características por recorte, agrupadas en siete familias (geometría y archivo, exposición y calidad, color global, histogramas HSV, textura LBP, forma HOG y variables construidas). Como las fotos de cada clase se tomaron en ráfagas, la partición es por bloques temporales: el 30 % final de las capturas de cada clase es prueba (799 recortes) y el resto es entrenamiento (2 088), con las 241 clases en ambos lados. La validación cruzada usa <code>StratifiedGroupKFold(3)</code> con tres bloques consecutivos por clase como grupos. Dentro de un <code>Pipeline</code> ajustado solo con entrenamiento se aplican: Box-Cox o Yeo-Johnson a las 13 variables con |asimetría| > 1, estandarización, el filtro de varianza (eliminó 1 variable) y el de redundancia con |ρ de Spearman| > 0.95 (eliminó 205). Quedan 616 candidatas.")

h2("s2", "2. Algoritmo baseline")
h3("s21", "2.1 Referencia de azar")
p("Antes de proponer un algoritmo se midió qué se obtiene sin aprender de las imágenes, con tres clasificadores triviales (Tabla 2). Con 241 clases, la exactitud esperada al azar uniforme es 1/241 ≈ 0.004 y ningún clasificador trivial supera un F1 macro de 0.006 en prueba. Cualquier modelo útil debe quedar muy por encima de ese valor.")
tabla("t2", "Clasificadores de referencia (azar)", ["Referencia", "F1 macro CV", "F1 macro prueba", "Exactitud prueba", "Top-3 prueba"], [
    ["Clase más frecuente", "0.0001", "0.0001", "0.0113", "0.0188"],
    ["Estratificado", "0.0037", "0.0034", "0.0038", "0.0113"],
    ["Uniforme", "0.0055", "0.0061", "0.0088", "0.0113"],
])
h3("s22", "2.2 Elección y justificación")
p("Se eligió una <b>regresión logística multinomial</b> sobre las características preparadas y seleccionadas con ANOVA. La elección responde a las características del problema:")
ul([
    "<b>Tipo de datos.</b> Tras la extracción, el problema es tabular: un vector numérico de longitud fija, con variables continuas y estandarizadas, para el que un modelo lineal es el punto de partida natural (Géron, 2022).",
    "<b>Cantidad de datos.</b> Hay entre 5 y 32 imágenes por clase. Un modelo con pocos hiperparámetros y regularización L2 tiene menos riesgo de memorizar que uno flexible, y entrena en segundos en CPU.",
    "<b>Características relevantes.</b> Como la regresión logística es sensible a las variables irrelevantes y a la escala, se combina con la preparación y la selección ANOVA del Avance 2. A cambio, sus coeficientes muestran qué variables usa.",
    "<b>Interpretabilidad y probabilidades.</b> Entrega una probabilidad por clase: una predicción de baja confianza puede mandarse a revisión manual en lugar de generar una alerta.",
    "<b>Continuidad.</b> Es el modelo con el que se validó la preparación del Avance 2.",
])
p("El número de características se ajustó con validación cruzada (Figura 2). El máximo está en <b>k = 50</b> (F1 macro CV 0.849 ± 0.014), el mismo valor del Avance 2. Con más variables el desempeño baja de forma sostenida (0.675 con las 616 candidatas), lo que indica sobreajuste cuando hay más dimensiones que información útil por clase.")
fig("f2", "celda-016", "F1 macro en validación cruzada contra el número de características seleccionadas por ANOVA.")
h3("s23", "2.3 Contraste con otros algoritmos")
p("Para comprobar que la elección no deja desempeño importante sobre la mesa, se contrastó con otros algoritmos simples bajo la misma preparación y validación (Tabla 3, Figura 3).")
tabla("t3", "Algoritmos con la misma preparación (k = 50)", ["Modelo", "F1 CV", "F1 prueba", "Exactitud prueba", "Top-3 prueba", "Segundos"], [
    ["<b>Regresión logística (baseline)</b>", "<b>0.849 ± 0.014</b>", "<b>0.881</b>", "<b>0.895</b>", "<b>0.951</b>", "<b>12.7</b>"],
    ["kNN (k = 3)", "0.775 ± 0.010", "0.814", "0.830", "0.922", "0.5"],
    ["Bayes ingenuo gaussiano", "0.685 ± 0.013", "0.817", "0.827", "0.874", "0.5"],
    ["Árbol de decisión", "0.513 ± 0.016", "0.606", "0.628", "0.628", "1.4"],
    ["Bosque aleatorio (300 árboles)", "0.867 ± 0.015", "0.901", "0.912", "0.966", "4.4"],
])
fig("f3", "celda-017", "F1 macro en validación cruzada (barras) y en prueba (puntos) por algoritmo.")
p("Todos los algoritmos superan ampliamente al azar. El bosque aleatorio obtiene el mejor F1, pero su ventaja en validación cruzada (0.018) es del orden de una desviación estándar entre pliegues. Se mantiene la regresión logística como baseline porque sus probabilidades están mejor calibradas (log-loss de 0.63 en prueba contra 0.99 del bosque), porque es interpretable y porque es el modelo de referencia del Avance 2. El bosque aleatorio queda registrado como el primer candidato a superar el baseline.")

h2("s3", "3. Importancia de las características")
p("Se combinaron tres tipos de método: filtro (estadístico F de ANOVA), embebido (magnitud media de los coeficientes de la regresión logística e impureza del bosque aleatorio) y permutación por familia sobre prueba. Se permuta por familia porque muchas variables están correlacionadas entre sí y permutarlas de una en una subestimaría su aporte. Además se hizo una ablación retirando cada familia.")
fig("f4", "celda-020", "Veinte variables más importantes según ANOVA F, coeficientes de la regresión logística e impureza del bosque aleatorio, por familia.")
tabla("t4", "Importancia por familia de características", ["Familia", "Candidatas", "Seleccionadas", "Caída de F1 al permutar", "Δ F1 al retirarla"], [
    ["Histogramas HSV", "20", "18", "0.562", "−0.106"],
    ["Color global", "13", "13", "0.424", "−0.058"],
    ["Textura LBP", "10", "8", "0.272", "−0.047"],
    ["Exposición y calidad", "7", "6", "0.184", "−0.032"],
    ["Geometría y archivo", "5", "4", "0.137", "−0.013"],
    ["Construidas", "7", "1", "0.014", "−0.005"],
    ["Forma HOG", "554", "0", "0.000", "0.000"],
])
fig("f5", "celda-021", "Importancia por permutación (prueba, 10 repeticiones) y ablación por familia.")
p("Los tres métodos coinciden en que <b>el color es la señal dominante</b>. Las variables con mayor F son el componente rojo-verde de Lab, los primeros intervalos de matiz, la saturación y el <i>colorfulness</i>: el color del empaque distingue marcas y presentaciones. La forma HOG no aporta en este modelo. Ninguna de sus variables entra en la selección, porque cada celda separa poco por sí sola y ANOVA evalúa las variables una por una; un modelo que combine los gradientes podría aprovecharla. La familia de geometría y archivo tiene importancia por permutación (0.14), pero retirarla casi no cambia el resultado (−0.013). Como las dos campañas de captura difieren en dispositivo y recorte, parte de esa señal puede deberse a cómo se tomó la foto y no al producto, por lo que conviene no depender de ella hasta tener capturas de varias sesiones por producto.")

h2("s4", "4. Sub y sobreajuste")
tabla("t5", "Métricas por conjunto", ["Modelo", "Conjunto", "F1 macro", "Exactitud", "Top-3", "Log-loss"], [
    ["Regresión logística", "Entrenamiento", "0.996", "0.997", "1.000", "0.207"],
    ["Regresión logística", "Validación cruzada", "0.849", "0.866", "0.943", "0.735"],
    ["Regresión logística", "Prueba", "0.881", "0.895", "0.951", "0.629"],
    ["Bosque aleatorio", "Entrenamiento", "1.000", "1.000", "1.000", "0.191"],
    ["Bosque aleatorio", "Validación cruzada", "0.867", "0.878", "0.950", "1.190"],
    ["Bosque aleatorio", "Prueba", "0.901", "0.912", "0.966", "0.992"],
])
p("El baseline muestra <b>sobreajuste moderado</b> (Tabla 5). El F1 macro es casi perfecto en entrenamiento (0.996) y baja a 0.849 en validación cruzada y a 0.881 en prueba, una brecha de unos 0.15 puntos; el log-loss sube de 0.21 a 0.74. No hay subajuste: el desempeño fuera de entrenamiento es alto y muy superior al azar. El bosque aleatorio memoriza por completo el entrenamiento, como es habitual en ensambles de árboles profundos.")
fig("f6", "celda-025", "Curva de aprendizaje y curva de validación sobre la regularización C del baseline.")
p("La <b>curva de aprendizaje</b> (Figura 6, izquierda) muestra que el F1 de validación sube sin estabilizarse (de 0.42 a 0.85 al pasar de 278 a 1 391 recortes) mientras el de entrenamiento se mantiene cerca de 1. Es el patrón de alta varianza: el modelo mejoraría con más imágenes por clase. La <b>curva de validación</b> (derecha) muestra subajuste con regularización fuerte (C ≤ 0.01, F1 de 0.02 a 0.22) y una meseta entre C = 1 y C = 10 (0.849 y 0.854). El valor por defecto ya está en la zona estable; su ajuste fino corresponde a la optimización de hiperparámetros. Que la prueba supere a la validación cruzada es esperable, porque el modelo final entrena con todo el conjunto de entrenamiento y no con dos tercios.")

h2("s5", "5. Métrica adecuada para el problema")
p("Si el modelo reconoce el producto equivocado, una posición correcta se reporta como mal ubicada (una alerta falsa que cuesta tiempo de personal) o un producto fuera de lugar pasa como correcto (un incumplimiento no detectado). El costo es similar para cualquier producto, y en un planograma cada producto ocupa pocas posiciones. Por eso la <b>métrica principal es el F1 macro</b>: promedia el F1 de cada clase con el mismo peso y combina la precisión (evitar alertas falsas) con la exhaustividad (no dejar pasar productos), sin dejarse dominar por las clases con más imágenes (Sokolova y Lapalme, 2009). Se complementa con la exactitud balanceada, la proporción de clases con exhaustividad baja y la exactitud top-3, útil si la operación incluye una revisión asistida.")
tabla("t6", "Desempeño del baseline en prueba por tipo de carpeta", ["Tipo de carpeta", "Clases", "F1 macro", "Precisión media", "Exhaustividad media", "Clases con exhaustividad < 0.5"], [
    ["Con código", "35", "0.964", "0.971", "0.967", "0"],
    ["Agrupamiento", "206", "0.866", "0.887", "0.882", "17"],
    ["<b>Total</b>", "<b>241</b>", "<b>0.881</b>", "<b>0.899</b>", "<b>0.895</b>", "<b>17</b>"],
])
fig("f7", "celda-028", "F1 por clase y relación entre precisión y exhaustividad por clase en prueba.")
p("El top-3 en prueba es de 0.951: una revisión que muestre tres alternativas recuperaría casi todos los errores. El desglose por tipo de carpeta es el hallazgo principal (Tabla 6). Las clases con código alcanzan 0.964, coherente con el Avance 2 aun compitiendo contra 241 clases. Todos los errores graves están en las carpetas de agrupamiento: las ocho clases con F1 = 0 y los diez pares más confundidos. Varios de esos pares tienen números consecutivos (por ejemplo, GRP_198 y GRP_199), lo que sugiere que algunos agrupamientos son el mismo producto o variantes casi idénticas separadas en dos carpetas. En ese caso el error estaría en la etiqueta y no en el modelo. <b>Validar con Bnext la correspondencia de los agrupamientos con productos reales es la acción con mayor impacto esperado en la métrica.</b> La exactitud global (0.895) ocultaría este problema; el F1 macro y el desglose por clase lo hacen visible.")

h2("s6", "6. Desempeño mínimo")
p("Bnext no ha proporcionado una medición histórica del proceso actual de revisión del anaquel, así que no existe un valor histórico de comparación. El desempeño mínimo se estableció con dos referencias: un <b>piso estadístico</b>, el mejor clasificador de azar, que el baseline debe superar fuera del intervalo de confianza del 95 % calculado por bootstrap con 1 000 remuestreos de prueba (Efron y Tibshirani, 1993); y un <b>piso de negocio</b>, la meta del Avance 0 de exactitud ≥ 0.85 en el estado asignado a cada posición. Como ese estado depende de reconocer el producto, se adopta F1 macro ≥ 0.85 para el reconocimiento. Es una meta propuesta por el equipo, pendiente de validar con el líder técnico de Bnext.")
tabla("t7", "Baseline frente a los pisos de desempeño (prueba)", ["Modelo", "F1 macro", "IC 95 %", "Supera el azar", "Meta 0.85 (puntual)", "Meta 0.85 (IC inferior)"], [
    ["<b>Baseline</b>", "<b>0.881</b>", "0.842 – 0.887", "Sí", "Sí", "No"],
    ["Azar (uniforme)", "0.006", "0.002 – 0.010", "-", "No", "No"],
    ["Baseline, clases con código", "0.964", "-", "-", "Sí", "-"],
    ["Baseline, clases de agrupamiento", "0.866", "-", "-", "Sí", "-"],
])
fig("f8", "celda-032", "Distribución bootstrap del F1 macro del baseline y del azar frente a la meta de negocio.")
p("El baseline supera el piso estadístico por más de dos órdenes de magnitud: los datos contienen información suficiente para reconocer el producto y el problema es viable con métodos simples. Frente al piso de negocio, cumple la meta por su valor puntual, pero el extremo inferior del intervalo (0.842) queda ligeramente por debajo. El intervalo es asimétrico porque, con pocas imágenes de prueba por clase, el remuestreo deja fuera algunas clases y tiende a subestimar el F1 macro. La conclusión prudente es que el baseline <b>está en el límite de la meta</b>. <b>El desempeño mínimo para los modelos siguientes queda fijado en F1 macro 0.881 y top-3 0.951 sobre esta misma partición de prueba.</b>")

h2("s7", "7. Conclusiones de la fase de Ingeniería del modelo")
ol([
    "<b>Algoritmo.</b> Regresión logística multinomial sobre 50 características seleccionadas con ANOVA. Se justifica por el tipo de datos, la poca cantidad de imágenes por clase, la interpretabilidad y las probabilidades calibradas. Entrena en unos 13 segundos en CPU.",
    "<b>Características.</b> El color explica la mayor parte del desempeño, seguido de la textura LBP y la exposición. HOG no aporta con selección univariada, y la geometría del archivo debe vigilarse como posible artefacto de captura.",
    "<b>Ajuste.</b> Sobreajuste moderado (0.996 contra 0.849). La curva de aprendizaje sigue subiendo, así que más imágenes por clase es la palanca principal.",
    "<b>Métrica.</b> F1 macro, con top-3 y desglose por clase; los errores se concentran en las carpetas de agrupamiento.",
    "<b>Desempeño mínimo.</b> El baseline (0.881) supera al azar (0.006) y alcanza en el límite la meta de 0.85; queda como el piso de los modelos siguientes.",
])
p("<b>Siguientes pasos.</b> Validar con Bnext la correspondencia de los agrupamientos con productos, y la posible fusión de carpetas. Solicitar capturas de más de una sesión por producto, para medir generalización y no solo separabilidad. Anotar un subconjunto de escenas para poder evaluar la detección en el anaquel completo. En el Avance 4, comparar contra este baseline el bosque aleatorio y modelos con representaciones aprendidas, con optimización de hiperparámetros.")

BIBLIO = [
    "Dalal, N., y Triggs, B. (2005). Histograms of oriented gradients for human detection. En <i>2005 IEEE Computer Society Conference on Computer Vision and Pattern Recognition</i> (Vol. 1, pp. 886–893). IEEE. https://doi.org/10.1109/CVPR.2005.177",
    "Efron, B., y Tibshirani, R. J. (1993). <i>An introduction to the bootstrap</i>. Chapman & Hall.",
    "Géron, A. (2022). <i>Hands-on machine learning with Scikit-Learn, Keras, and TensorFlow</i> (3a ed.). O'Reilly Media.",
    "Ojala, T., Pietikäinen, M., y Mäenpää, T. (2002). Multiresolution gray-scale and rotation invariant texture classification with local binary patterns. <i>IEEE Transactions on Pattern Analysis and Machine Intelligence, 24</i>(7), 971–987. https://doi.org/10.1109/TPAMI.2002.1017623",
    "Pedregosa, F., Varoquaux, G., Gramfort, A., Michel, V., Thirion, B., Grisel, O., Blondel, M., Prettenhofer, P., Weiss, R., Dubourg, V., Vanderplas, J., Passos, A., Cournapeau, D., Brucher, M., Perrot, M., y Duchesnay, É. (2011). Scikit-learn: Machine learning in Python. <i>Journal of Machine Learning Research, 12</i>, 2825–2830.",
    "Sokolova, M., y Lapalme, G. (2009). A systematic analysis of performance measures for classification tasks. <i>Information Processing &amp; Management, 45</i>(4), 427–437. https://doi.org/10.1016/j.ipm.2009.03.002",
    "Visengeriyeva, L., Kammer, A., Bär, I., Kniesz, A., y Plöd, M. (2023). <i>CRISP-ML(Q). The ML lifecycle process</i>. MLOps. INNOQ. https://ml-ops.org/content/crisp-ml",
]


INTEGRANTES = [("Jesús Soto Medrano", "A01794746"), ("Germán Eduardo Rodríguez Vázquez", "A01796325"),
               ("Oscar Alberto Ramírez Anaya", "A01795438")]
TITULO = "Planograma: IA para detección de productos y ubicación"
SUBTITULO = "Modelo de referencia para el reconocimiento de producto en recortes de anaquel: algoritmo, características, ajuste, métrica y desempeño mínimo"


# ------------------------------------------------------------------ HTML
def generar_html(figs, paginas):
    def pag(ref):
        return paginas.get(ref, "")
    tablas = [b for b in B if b[0] == "tabla"]
    figuras = [b for b in B if b[0] == "fig"]
    num_t = {b[1]: i + 1 for i, b in enumerate(tablas)}
    num_f = {b[1]: i + 1 for i, b in enumerate(figuras)}
    logo = base64.b64encode(LOGO.read_bytes()).decode()
    o = ['<!DOCTYPE html>\n<html lang="es">\n<head>\n<meta charset="utf-8">',
         "<title>Avance 3. Baseline - Equipo 48</title>",
         '<link rel="stylesheet" href="assets/print.css">',
         "<style>.logo-tec{width:62mm;margin:0 auto 4mm auto;display:block}"
         "figure img{max-width:100%;max-height:150mm}"
         "code{font-family:Menlo,Consolas,monospace;font-size:9pt;hyphens:none}</style>",
         "</head>\n<body>",
         '<section class="portada">\n<div class="portada-encabezado">',
         f'<img class="logo-tec" alt="Tecnológico de Monterrey" src="data:image/png;base64,{logo}">',
         '<p class="programa">Maestría en Inteligencia Artificial Aplicada</p>',
         '<p class="curso">TC5035.10 — Proyecto Integrador (Grupo 10) · Trimestre septiembre – diciembre 2026</p>',
         '<hr class="regla">\n</div>\n<div>',
         '<p class="etiqueta">Avance 3 · Baseline</p>',
         f'<h1 class="titulo">{TITULO}</h1>', f'<p class="subtitulo">{SUBTITULO}</p>',
         '</div>\n<div>\n<p class="equipo">Equipo 48</p>\n<ul class="integrantes">']
    o += [f'<li><span>{n}</span><span class="matricula">{m}</span></li>' for n, m in INTEGRANTES]
    o += ["</ul>\n</div>", '<div class="portada-pie">',
          "<p><strong>Profesora:</strong> Dra. Grettel Barceló Alonso</p>",
          "<p><strong>Asesor:</strong> Dr. Raúl Valente Ramírez</p>",
          "<p><strong>Sponsor:</strong> Dr. Juan Arturo Nolazco Flores, Director del Hub de Ciencia de Datos, Tecnológico de Monterrey</p>",
          "<p><strong>Organización:</strong> Bnext, Monterrey, Nuevo León</p>",
          f'<p style="margin-top:4mm">{FECHA}</p>', "</div>\n</section>"]
    o.append('<h2>Índice</h2>\n<ul class="indice">')
    for b in B:
        if b[0] in ("h2", "h3"):
            cls = ' class="sub"' if b[0] == "h3" else ""
            o.append(f'<li{cls}><span>{b[2]}</span><span class="relleno"></span><span class="pag">{pag(b[1])}</span></li>')
    o.append(f'<li><span>Bibliografía</span><span class="relleno"></span><span class="pag">{pag("biblio")}</span></li></ul>')
    o.append('<h3 class="lista-indice-titulo">Índice de tablas</h3>\n<ul class="indice">')
    for b in tablas:
        o.append(f'<li class="sub"><span>Tabla {num_t[b[1]]}. {b[2]}</span><span class="relleno"></span><span class="pag">{pag(b[1])}</span></li>')
    o.append('</ul>\n<h3 class="lista-indice-titulo">Índice de figuras</h3>\n<ul class="indice">')
    for b in figuras:
        corto = b[3].split(". ")[0].rstrip(".")
        o.append(f'<li class="sub"><span>Figura {num_f[b[1]]}. {corto}</span><span class="relleno"></span><span class="pag">{pag(b[1])}</span></li>')
    o.append("</ul>")
    primero = True
    for b in B:
        t = b[0]
        if t == "h2":
            salto = ' class="salto"' if primero else ""
            o.append(f'<h2{salto} id="{b[1]}">{b[2]}</h2>'); primero = False
        elif t == "h3":
            o.append(f'<h3 id="{b[1]}">{b[2]}</h3>')
        elif t == "p":
            o.append(f"<p>{b[1]}</p>")
        elif t in ("ul", "ol"):
            o.append(f"<{t}>" + "".join(f"<li>{x}</li>" for x in b[1]) + f"</{t}>")
        elif t == "nota":
            o.append(f'<div class="nota"><p class="titulo-nota">{b[1]}</p><p>{b[2]}</p></div>')
        elif t == "tabla":
            _, i, tit, enc, filas = b
            o.append(f'<table id="{i}"><caption><b>Tabla {num_t[i]}.</b> {tit}</caption><thead><tr>'
                     + "".join(f"<th>{e}</th>" for e in enc) + "</tr></thead><tbody>"
                     + "".join("<tr>" + "".join(f"<td>{c}</td>" for c in f) + "</tr>" for f in filas)
                     + "</tbody></table>")
        elif t == "fig":
            _, i, celda, tit = b
            o.append(f'<figure id="{i}"><img alt="Figura {num_f[i]}" src="data:image/png;base64,{figs[celda]}">'
                     f"<figcaption><b>Figura {num_f[i]}.</b> {tit} Elaboración propia con datos de Bnext.</figcaption></figure>")
    o.append('<h2 id="biblio">Bibliografía</h2>\n<div class="biblio">')
    o += [f"<p>{x}</p>" for x in BIBLIO]
    o.append("</div>\n</body>\n</html>\n")
    HTML.write_text("\n".join(o), encoding="utf-8")


def imprimir_pdf():
    subprocess.run([CHROME, "--headless", "--disable-gpu", "--no-pdf-header-footer",
                    f"--print-to-pdf={PDF}", HTML.as_uri()], check=True, capture_output=True)


def localizar_paginas():
    """Página (1-based) donde aparece cada encabezado, tabla y figura, después del índice."""
    doc = fitz.open(PDF)
    textos = [pg.get_text() for pg in doc]
    paginas = {}
    tablas = [b for b in B if b[0] == "tabla"]
    figuras = [b for b in B if b[0] == "fig"]
    # El índice repite los títulos: la búsqueda empieza después de su última entrada
    ultima = f"Figura {len(figuras)}."
    desde = next(n for n in range(2, len(textos) + 1) if ultima in textos[n - 1]) + 1
    num_t = {b[1]: i + 1 for i, b in enumerate(tablas)}
    num_f = {b[1]: i + 1 for i, b in enumerate(figuras)}
    def buscar(clave):
        for n in range(desde, len(textos) + 1):
            if clave in textos[n - 1]:
                return n
        return ""
    for b in B:
        if b[0] in ("h2", "h3"):
            paginas[b[1]] = buscar(b[2])
        elif b[0] == "tabla":
            paginas[b[1]] = buscar(f"Tabla {num_t[b[1]]}.")
        elif b[0] == "fig":
            paginas[b[1]] = buscar(f"Figura {num_f[b[1]]}.")
    paginas["biblio"] = buscar("Bibliografía")
    return paginas


def pie_pdf():
    doc = fitz.open(PDF)
    total = doc.page_count
    for i, pg in enumerate(doc, start=1):
        if i == 1:
            continue
        w, h = pg.rect.width, pg.rect.height
        y = h - 42.5
        pg.draw_line(fitz.Point(56.7, y - 9), fitz.Point(w - 56.7, y - 9), color=(0.79, 0.81, 0.84), width=0.5)
        pg.insert_text(fitz.Point(56.7, y + 2), "TC5035.10 Proyecto Integrador  ·  Avance 3  ·  Equipo 48",
                       fontname="helv", fontsize=7.5, color=(0.42, 0.42, 0.42))
        der = f"Página {i} de {total}"
        pg.insert_text(fitz.Point(w - 56.7 - fitz.get_text_length(der, fontname="helv", fontsize=7.5), y + 2),
                       der, fontname="helv", fontsize=7.5, color=(0.42, 0.42, 0.42))
    doc.set_metadata({"title": "Avance 3. Baseline — Equipo 48", "author": "Equipo 48",
                      "subject": "TC5035.10 Proyecto Integrador — Avance 3",
                      "keywords": "planograma, baseline, regresion logistica, importancia de caracteristicas, F1 macro, CRISP-ML(Q)"})
    doc.saveIncr()
    return total


# ------------------------------------------------------------------ DOCX
def runs_desde_html(par, texto, tam=None, color=None, fuente=None):
    """Convierte el marcado mínimo (<b>, <i>, <code>) a runs de Word."""
    estado = {"b": False, "i": False, "code": False}
    for trozo in re.split(r"(</?(?:b|i|code)>)", texto):
        m = re.fullmatch(r"<(/?)(b|i|code)>", trozo)
        if m:
            estado[m.group(2)] = m.group(1) == ""
            continue
        if not trozo:
            continue
        r = par.add_run(html.unescape(trozo))
        r.bold = estado["b"] or None
        r.italic = estado["i"] or None
        if estado["code"]:
            r.font.name = "Consolas"; r.font.size = Pt(9)
        elif fuente:
            r.font.name = fuente
        if tam and not estado["code"]:
            r.font.size = Pt(tam)
        if color:
            r.font.color.rgb = color


def sombrear(celda, hexcolor):
    tcPr = celda._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hexcolor)
    tcPr.append(shd)


def campo(par, instruccion):
    for tipo, texto in (("begin", None), (None, instruccion), ("separate", None), (None, "1"), ("end", None)):
        r = par.add_run()
        if tipo:
            fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), tipo); r._r.append(fc)
        elif texto == instruccion:
            it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = instruccion; r._r.append(it)
        else:
            r.text = texto


def generar_docx(figs, paginas):
    d = Document()
    sec = d.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.0)
    sec.top_margin, sec.bottom_margin = Cm(2.2), Cm(2.4)
    normal = d.styles["Normal"]
    normal.font.name = "Georgia"; normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6); normal.paragraph_format.line_spacing = 1.15
    for nombre, tam in (("Heading 1", 15), ("Heading 2", 12)):
        st = d.styles[nombre]
        st.font.name = "Arial"; st.font.size = Pt(tam); st.font.color.rgb = AZUL; st.font.bold = True
        rpr = st.element.get_or_add_rPr(); rf = rpr.find(qn("w:rFonts"))
        for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            rf.set(qn(a), "Arial")
    def centrado(texto, tam=None, color=None, fuente=None, negrita=False, antes=0):
        par = d.add_paragraph(); par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.space_before = Pt(antes)
        runs_desde_html(par, f"<b>{texto}</b>" if negrita else texto, tam, color, fuente)
        return par
    # portada
    par = d.add_paragraph(); par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    par.add_run().add_picture(str(LOGO), width=Cm(6.2))
    centrado("Maestría en Inteligencia Artificial Aplicada", 11, RGBColor(0x44, 0x44, 0x44), antes=30)
    centrado("TC5035.10 — Proyecto Integrador (Grupo 10) · Trimestre septiembre – diciembre 2026", 10, RGBColor(0x44, 0x44, 0x44))
    centrado("AVANCE 3 · BASELINE", 8.5, RGBColor(0x66, 0x66, 0x66), "Arial", antes=40)
    centrado(TITULO, 20, AZUL, "Arial", negrita=True, antes=6)
    centrado(f"<i>{SUBTITULO}</i>", 10.5, RGBColor(0x55, 0x55, 0x55))
    centrado("Equipo 48", 13, None, "Arial", negrita=True, antes=40)
    t = d.add_table(rows=0, cols=2); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for n, m in INTEGRANTES:
        fila = t.add_row().cells
        fila[0].text = n; fila[1].text = m
        fila[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    centrado("<b>Profesora:</b> Dra. Grettel Barceló Alonso", 9.5, antes=50)
    centrado("<b>Asesor:</b> Dr. Raúl Valente Ramírez", 9.5)
    centrado("<b>Sponsor:</b> Dr. Juan Arturo Nolazco Flores, Director del Hub de Ciencia de Datos, Tecnológico de Monterrey", 9.5)
    centrado("<b>Organización:</b> Bnext, Monterrey, Nuevo León", 9.5)
    centrado(FECHA, 9.5, RGBColor(0x66, 0x66, 0x66), antes=8)
    # índice (campo TOC; Word lo actualiza al abrir)
    d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    par = d.add_paragraph(); runs_desde_html(par, "Índice", 15, AZUL, "Arial")
    campo(d.add_paragraph(), 'TOC \\o "1-2" \\h \\z \\u')
    ajustes = d.settings.element
    uf = OxmlElement("w:updateFields"); uf.set(qn("w:val"), "true"); ajustes.append(uf)
    d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    # cuerpo
    tablas = [b for b in B if b[0] == "tabla"]; figuras = [b for b in B if b[0] == "fig"]
    num_t = {b[1]: i + 1 for i, b in enumerate(tablas)}; num_f = {b[1]: i + 1 for i, b in enumerate(figuras)}
    for b in B:
        tp = b[0]
        if tp == "h2":
            d.add_heading(b[2], level=1)
        elif tp == "h3":
            d.add_heading(b[2], level=2)
        elif tp == "p":
            par = d.add_paragraph(); par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY; runs_desde_html(par, b[1])
        elif tp in ("ul", "ol"):
            for x in b[1]:
                par = d.add_paragraph(style="List Number" if tp == "ol" else "List Bullet"); runs_desde_html(par, x)
        elif tp == "nota":
            tn = d.add_table(rows=1, cols=1); c = tn.rows[0].cells[0]; sombrear(c, "F3F6FB")
            runs_desde_html(c.paragraphs[0], f"<b>{b[1]}</b>", 9.5, AZUL, "Arial")
            runs_desde_html(c.add_paragraph(), b[2], 9.5)
            d.add_paragraph()
        elif tp == "tabla":
            _, i, tit, enc, filas = b
            par = d.add_paragraph(); par.paragraph_format.keep_with_next = True
            runs_desde_html(par, f"<b>Tabla {num_t[i]}.</b> {tit}", 9, AZUL, "Arial")
            tb = d.add_table(rows=1, cols=len(enc)); tb.style = "Table Grid"
            for j, e in enumerate(enc):
                c = tb.rows[0].cells[j]; sombrear(c, AZUL_TABLA)
                runs_desde_html(c.paragraphs[0], f"<b>{e}</b>", 9, RGBColor(0xFF, 0xFF, 0xFF), "Arial")
            for f in filas:
                cs = tb.add_row().cells
                for j, v in enumerate(f):
                    runs_desde_html(cs[j].paragraphs[0], v, 9)
            d.add_paragraph()
        elif tp == "fig":
            _, i, celda, tit = b
            par = d.add_paragraph(); par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            par.paragraph_format.keep_with_next = True
            par.add_run().add_picture(io.BytesIO(base64.b64decode(figs[celda])), width=Cm(16))
            par = d.add_paragraph()
            runs_desde_html(par, f"<b>Figura {num_f[i]}.</b> {tit} Elaboración propia con datos de Bnext.", 9, None, "Arial")
    d.add_heading("Bibliografía", level=1)
    for x in BIBLIO:
        par = d.add_paragraph(); par.paragraph_format.left_indent = Cm(1); par.paragraph_format.first_line_indent = Cm(-1)
        runs_desde_html(par, x, 10)
    # pie con número de página
    sec.different_first_page_header_footer = True
    pie = sec.footer.paragraphs[0]
    runs_desde_html(pie, "TC5035.10 Proyecto Integrador  ·  Avance 3  ·  Equipo 48        Página ", 7.5, RGBColor(0x6B, 0x6B, 0x6B), "Arial")
    campo(pie, "PAGE")
    d.core_properties.author = "Equipo 48"
    d.core_properties.title = "Avance 3. Baseline — Equipo 48"
    d.save(DOCX)


if __name__ == "__main__":
    figs = figuras_libreta()
    generar_html(figs, {})
    imprimir_pdf()
    paginas = localizar_paginas()
    generar_html(figs, paginas)
    imprimir_pdf()
    assert localizar_paginas() == paginas, "La paginación cambió al insertar los números del índice"
    total = pie_pdf()
    generar_docx(figs, paginas)
    print(f"{PDF.name}: {total} páginas | {HTML.name} | {DOCX.name}")
